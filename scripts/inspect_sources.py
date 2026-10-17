# -*- coding: utf-8 -*-
"""探查三个源文件的表结构：sheet 名、行数、列名、前几行样本"""
import sys, io
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl, xlrd

D = r'C:\Users\Administrator\Desktop'
FILES = [
    ('车队三楼库存清单', D + r'\周转物资-车队三楼库存清单(1).xls', 'xls'),
    ('资产信息查询',     D + r'\资产信息查询(1).xlsx', 'xlsx'),
    ('李师傅周转物资明细', D + r'\李师傅周转物资明细(1).xlsx', 'xlsx'),
]

def show_xlsx(path, tag, maxsheet=6):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print('=' * 70)
    print(f'【{tag}】sheets = {wb.sheetnames}')
    for name in wb.sheetnames[:maxsheet]:
        ws = wb[name]
        print(f'\n--- sheet: {name}  dims={ws.max_row}行 x {ws.max_column}列 ---')
        for i, row in enumerate(ws.iter_rows(max_row=6, values_only=True)):
            cells = ['' if c is None else str(c)[:24] for c in row[:16]]
            print(f'  r{i}: ' + ' | '.join(cells))
    wb.close()

def show_xls(path, tag):
    book = xlrd.open_workbook(path)
    print('=' * 70)
    print(f'【{tag}】sheets = {book.sheet_names()}')
    for name in book.sheet_names()[:6]:
        sh = book.sheet_by_name(name)
        print(f'\n--- sheet: {name}  {sh.nrows}行 x {sh.ncols}列 ---')
        for r in range(min(6, sh.nrows)):
            cells = [str(sh.cell_value(r, c))[:24] for c in range(min(16, sh.ncols))]
            print(f'  r{r}: ' + ' | '.join(cells))

for tag, path, kind in FILES:
    try:
        if kind == 'xls':
            show_xls(path, tag)
        else:
            show_xlsx(path, tag)
    except Exception as e:
        print(f'[{tag}] 读取失败: {type(e).__name__}: {e}')
    print()